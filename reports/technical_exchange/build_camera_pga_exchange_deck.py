from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from pptx.enum.text import PP_ALIGN

from build_technical_exchange_deck_v5 import (
    ASSET_DIR,
    BG,
    BLUE,
    CYAN,
    GREEN,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    PANEL,
    PANEL_BLUE,
    PANEL_GREEN,
    PANEL_ORANGE,
    PANEL_RED,
    PURPLE,
    RED,
    SLIDE_H,
    SLIDE_W,
    TEXT,
    WHITE,
    Deck,
    add_bullets,
    add_card,
    add_image_fit,
    add_kpi,
    add_line,
    add_pill,
    add_shape,
    add_table,
    add_textbox,
)

OUT_DIR = Path(__file__).resolve().parent
PGA_ASSET_DIR = OUT_DIR.parent / "pga_academic_report_assets_current"
PPTX_PATH = OUT_DIR / "camera_pga_exchange_20260629.pptx"
CHAOSUAN_DIR = OUT_DIR.parents[2] / "chaosuan_res"


def tech_asset(name: str) -> Path:
    path = ASSET_DIR / name
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def pga_asset(name: str) -> Path:
    path = PGA_ASSET_DIR / name
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def add_arrow(slide, x1, y1, x2, y2, color=CYAN, width=1.8):
    line = add_line(slide, x1, y1, x2, y2, color, width)
    line.line.end_arrowhead = True
    return line


def _pga_metrics(npz_path: Path, split: str) -> dict[str, float]:
    z = np.load(npz_path, allow_pickle=True)
    pred = np.asarray(z[f"{split}_pga_mu_best"], dtype=float)
    label = np.asarray(z[f"{split}_pga_label"], dtype=float).squeeze(-1)
    valid = np.asarray(z[f"{split}_pga_target_valid"]).astype(bool)
    x = label[valid].ravel()
    y = pred[valid].ravel()
    err = y - x
    ss_res = float(np.sum(err**2))
    ss_tot = float(np.sum((x - x.mean()) ** 2))
    return {
        "mae": float(np.mean(np.abs(err))),
        "rmse": float(np.sqrt(np.mean(err**2))),
        "corr": float(np.corrcoef(x, y)[0, 1]),
        "r2": 1.0 - ss_res / ss_tot,
        "slope": float(np.polyfit(x, y, 1)[0]),
        "n": float(len(x)),
    }


def rt44_rt51_table() -> pd.DataFrame:
    runs = [
        ("rt44", "temporal residual scale0", "weights_japan_overfit_pga15_stage2_512_rt44_cached_dpk_event_temporal_residual_scale0"),
        ("rt45", "temporal residual scale2", "weights_japan_overfit_pga15_stage2_512_rt45_cached_dpk_event_temporal_residual_scale2"),
        ("rt46", "temporal residual scale4", "weights_japan_overfit_pga15_stage2_512_rt46_cached_dpk_event_temporal_residual_scale4"),
        ("rt48", "station-pool + temporal", "weights_japan_overfit_pga15_stage2_512_rt48_cached_dpk_event_station_pool_temporal_residual_scale4"),
        ("rt49", "layerwise temporal", "weights_japan_overfit_pga15_stage2_512_rt49_cached_dpk_event_layerwise_temporal_readout_scale4"),
        ("rt50", "independent residual", "weights_japan_overfit_pga15_stage2_512_rt50_cached_dpk_event_temporal_residual_independent_scale4"),
        ("rt51", "station-roll control", "weights_japan_overfit_pga15_stage2_512_rt51_cached_dpk_event_temporal_residual_stationroll_control_scale4"),
    ]
    rows = []
    for exp, method, dirname in runs:
        path = CHAOSUAN_DIR / dirname / "eval_results_last.npz"
        train = _pga_metrics(path, "train")
        test = _pga_metrics(path, "val")
        rows.append(
            {
                "Exp": exp,
                "Setting": method,
                "Train MAE": train["mae"],
                "Train Corr": train["corr"],
                "Train slope": train["slope"],
                "Test MAE*": test["mae"],
                "Test Corr*": test["corr"],
                "Test slope*": test["slope"],
            }
        )
    return pd.DataFrame(rows)


def _pga_metrics_at_time(npz_path: Path, split: str, elapsed_seconds: float) -> dict[str, float]:
    z = np.load(npz_path, allow_pickle=True)
    elapsed = np.asarray(z[f"{split}_realtime_elapsed_time"], dtype=float).reshape(-1)
    pred = np.asarray(z[f"{split}_pga_mu_best"], dtype=float)
    label = np.asarray(z[f"{split}_pga_label"], dtype=float).squeeze(-1)
    valid = np.asarray(z[f"{split}_pga_target_valid"]).astype(bool)
    row_mask = np.isclose(elapsed, elapsed_seconds)
    mask = valid & row_mask[:, None]
    x = label[mask].ravel()
    y = pred[mask].ravel()
    err = y - x
    return {
        "mae": float(np.mean(np.abs(err))),
        "corr": float(np.corrcoef(x, y)[0, 1]),
        "n": float(len(x)),
    }


def rt48_elapsed_detail_table() -> pd.DataFrame:
    path = CHAOSUAN_DIR / "weights_japan_overfit_pga15_stage2_512_rt48_cached_dpk_event_station_pool_temporal_residual_scale4" / "eval_results_last.npz"
    rows = []
    for t in [1, 3, 5, 10, 20, 40, 90]:
        train = _pga_metrics_at_time(path, "train", t)
        test = _pga_metrics_at_time(path, "val", t)
        rows.append(
            {
                "Time": f"{t}s",
                "Train MAE": train["mae"],
                "Train Corr": train["corr"],
                "Test MAE*": test["mae"],
                "Test Corr*": test["corr"],
            }
        )
    return pd.DataFrame(rows)


def rt_elapsed_test_mae_table() -> pd.DataFrame:
    runs = [
        ("rt44", "weights_japan_overfit_pga15_stage2_512_rt44_cached_dpk_event_temporal_residual_scale0"),
        ("rt45", "weights_japan_overfit_pga15_stage2_512_rt45_cached_dpk_event_temporal_residual_scale2"),
        ("rt46", "weights_japan_overfit_pga15_stage2_512_rt46_cached_dpk_event_temporal_residual_scale4"),
        ("rt48", "weights_japan_overfit_pga15_stage2_512_rt48_cached_dpk_event_station_pool_temporal_residual_scale4"),
        ("rt49", "weights_japan_overfit_pga15_stage2_512_rt49_cached_dpk_event_layerwise_temporal_readout_scale4"),
        ("rt50", "weights_japan_overfit_pga15_stage2_512_rt50_cached_dpk_event_temporal_residual_independent_scale4"),
        ("rt51", "weights_japan_overfit_pga15_stage2_512_rt51_cached_dpk_event_temporal_residual_stationroll_control_scale4"),
    ]
    rows = []
    for exp, dirname in runs:
        path = CHAOSUAN_DIR / dirname / "eval_results_last.npz"
        row = {"Exp": exp}
        for t in [1, 3, 5, 10, 20, 40, 90]:
            row[f"{t}s"] = _pga_metrics_at_time(path, "val", t)["mae"]
        rows.append(row)
    return pd.DataFrame(rows)


def title_slide(d: Deck):
    s = d.prs.slides.add_slide(d.blank)
    d.page_no += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    add_textbox(
        s,
        "基于视频摄像头与多台站波形的 PGA 估计",
        1.15,
        2.55,
        11.05,
        0.65,
        28,
        NAVY,
        True,
        PP_ALIGN.CENTER,
        margin=0.0,
    )
    add_textbox(s, "2026-06-29", 1.15, 3.42, 11.05, 0.30, 14, MUTED, False, PP_ALIGN.CENTER, margin=0.0)


def workflow_slide(d: Deck):
    s = d.slide("总体工作链条", "摄像头提供振动观测，多台站模型提供摄像头位置的参考 PGA", "Workflow", CYAN)
    steps = [
        ("视频摄像头", "地震时画面产生抖动", BLUE),
        ("振动曲线", "从视频运动中提取时间序列", CYAN),
        ("摄像头标定", "拟合位置相关乘数", ORANGE),
        ("PGA 估计", "输出摄像头位置 PGA", GREEN),
    ]
    x_positions = [0.70, 3.85, 7.00, 10.15]
    for i, ((title, body, color), x) in enumerate(zip(steps, x_positions)):
        add_card(s, x, 2.05, 2.45, 1.30, title, body, color, PANEL, number=str(i + 1))
        if i < len(steps) - 1:
            add_arrow(s, x + 2.50, 2.70, x + 3.05, 2.70, color)
    add_shape(s, 0.85, 4.65, 11.85, 0.88, PANEL_BLUE, LINE, radius=True)
    add_textbox(
        s,
        "深度学习 PGA 模型的作用：给历史地震中摄像头所在位置生成参考 PGA，用于拟合标定系数。",
        1.10,
        4.88,
        11.35,
        0.32,
        14.5,
        NAVY,
        True,
        PP_ALIGN.CENTER,
        margin=0.0,
    )


def task_split_slide(d: Deck):
    s = d.slide("当前任务拆解", "第一步已实现，当前重点是把相对振动强度转成绝对 PGA", "Task", BLUE)
    add_card(s, 0.85, 1.45, 5.70, 1.55, "1. 从视频抖动中提取振动曲线", "输入为地震期间摄像头视频；输出为描述画面运动的振动时间序列。", GREEN, PANEL, number="A")
    add_card(s, 6.85, 1.45, 5.70, 1.55, "2. 根据振动曲线估计 PGA", "需要建立视频振动强度和真实地面加速度之间的位置相关关系。", ORANGE, PANEL, number="B")
    add_shape(s, 1.10, 3.75, 11.25, 1.25, PANEL_ORANGE, LINE, radius=True)
    add_textbox(s, "核心转化", 1.35, 3.95, 1.30, 0.28, 12.5, ORANGE, True, margin=0.0)
    add_textbox(s, "视频振动曲线是相对运动观测；PGA 是地面运动绝对强度。两者之间需要摄像头位置的标定系数。", 2.55, 3.93, 9.35, 0.36, 14.0, TEXT, True, margin=0.0)
    add_pill(s, "类似地震仪的简化仪器响应因子，但对象变成每个摄像头位置", 2.50, 5.55, 8.30, 0.36, PANEL_BLUE, BLUE)


def calibration_factor_slide(d: Deck):
    s = d.slide("摄像头标定系数的定义", "用历史事件把视频振动强度映射到摄像头位置 PGA", "Calibration", ORANGE)
    add_shape(s, 0.85, 1.32, 11.70, 1.05, PANEL_BLUE, LINE, radius=True)
    add_textbox(s, "PGA_camera,event  =  k_camera  ×  Motion_video,event", 1.05, 1.62, 11.30, 0.38, 24, NAVY, True, PP_ALIGN.CENTER, margin=0.0)
    left = pd.DataFrame(
        [
            ["Motion_video,event", "视频振动曲线提取出的强度指标"],
            ["PGA_camera,event", "摄像头位置的参考 PGA"],
            ["k_camera", "该摄像头位置的标定系数"],
        ],
        columns=["符号", "含义"],
    )
    add_table(s, left, 1.00, 3.05, 5.25, 1.65, 10.0, header_color=ORANGE, widths=[1.7, 3.0])
    add_card(s, 6.85, 3.00, 5.35, 1.05, "拟合数据", "同一摄像头在多个历史事件中的视频振动强度与参考 PGA。", CYAN, PANEL)
    add_card(s, 6.85, 4.40, 5.35, 1.05, "在线使用", "地震发生后只需视频振动曲线和已标定系数，即可估计 PGA。", GREEN, PANEL)


def model_io_slide(d: Deck):
    s = d.slide("多台站 PGA 模型的输入输出", "给定若干地震台波形，查询任意目标位置 PGA", "Model", CYAN)
    add_card(s, 0.80, 1.35, 3.65, 1.25, "输入台站集合", "台站坐标、三分量波形、有效台站 mask", BLUE, PANEL, number="1")
    add_card(s, 4.85, 1.35, 3.65, 1.25, "目标位置查询", "可为强震台站，也可为摄像头位置", CYAN, PANEL, number="2")
    add_card(s, 8.90, 1.35, 3.65, 1.25, "输出", "目标位置 log PGA / PGA", GREEN, PANEL, number="3")
    add_image_fit(s, tech_asset("task_setup_multistation.png"), 1.05, 3.05, 11.10, 3.30)
    add_pill(s, "同一个事件可查询多个目标位置，因此适合为摄像头网络批量生成参考 PGA", 2.20, 6.55, 8.95, 0.34, PANEL_BLUE, BLUE)


def architecture_slide(d: Deck):
    s = d.slide("模型结构简述", "DiTing 提取单台站波形表征，TEAM-style 模块融合多台站信息", "Architecture", CYAN)
    add_image_fit(s, tech_asset("architecture_implementation_detail.png"), 0.50, 1.18, 12.35, 4.90)
    add_card(s, 0.82, 6.16, 2.75, 0.58, "DiTing encoder", "单台站波形表征", BLUE, PANEL_BLUE, title_size=10.8, body_size=8.5)
    add_card(s, 3.85, 6.16, 2.75, 0.58, "Station adapter", "转换为台站 token", CYAN, PANEL_BLUE, title_size=10.8, body_size=8.5)
    add_card(s, 6.88, 6.16, 2.75, 0.58, "TEAM Transformer", "多台站集合建模", ORANGE, PANEL_ORANGE, title_size=10.8, body_size=8.5)
    add_card(s, 9.92, 6.16, 2.75, 0.58, "Target readout", "输出目标 PGA", GREEN, PANEL_GREEN, title_size=10.8, body_size=8.5)


def data_training_slide(d: Deck):
    s = d.slide("数据与训练设置", "以日本 K-NET / KiK-net 强震数据训练和评估多台站 PGA 模型", "Data", BLUE)
    split = pd.DataFrame(
        [
            ["Train", 699, "30,506", "M 3.7", "28.28 m/s²"],
            ["Validation", 101, "8,544", "M 4.0", "5.72 m/s²"],
            ["Test", 201, "13,052", "M 4.0", "5.74 m/s²"],
        ],
        columns=["Split", "Events", "Station records", "Median M", "Max PGA"],
    )
    add_table(s, split, 0.75, 1.35, 6.30, 1.55, 9.8, header_color=BLUE, widths=[1.2, 1.1, 1.8, 1.1, 1.3])
    add_image_fit(s, pga_asset("data_split_summary.png"), 7.35, 1.22, 5.05, 2.40)
    add_image_fit(s, pga_asset("data_distribution_overview.png"), 0.80, 3.55, 5.85, 2.55)
    add_card(s, 7.20, 4.18, 5.10, 0.88, "评估口径", "主指标在 log PGA 空间计算，包括 MAE、RMSE、Corr、R²、slope。", CYAN, PANEL)
    add_card(s, 7.20, 5.36, 5.10, 0.88, "输入设置", "模型支持可变数量输入台站，并对目标位置逐点输出 PGA。", GREEN, PANEL)


def data_spatial_distribution_slide(d: Deck):
    s = d.slide("训练数据空间分布", "事件和强震台站覆盖日本主要地震活动区域", "Data", BLUE)
    add_image_fit(s, pga_asset("data_event_station_map.png"), 0.55, 1.20, 7.05, 5.45)
    add_card(s, 8.00, 1.45, 4.35, 1.05, "台站位置", "K-NET / KiK-net 提供分布式强震观测，是多台站 PGA 模型的输入基础。", BLUE, PANEL)
    add_card(s, 8.00, 2.88, 4.35, 1.05, "事件位置", "训练事件覆盖不同震中位置、震级和震源深度。", CYAN, PANEL)
    add_card(s, 8.00, 4.31, 4.35, 1.05, "目标位置", "评估时可把任意台站或摄像头位置作为 target query。", GREEN, PANEL)
    add_pill(s, "这张图对应深度模型训练数据分布，不代表摄像头布设位置", 8.10, 5.90, 4.15, 0.34, PANEL_ORANGE, ORANGE)


def results_slide(d: Deck):
    s = d.slide("已有全数据 PGA 估计结果", "multi-station full model 优于 single-station 参考模型", "Results", GREEN)
    metrics = pd.DataFrame(
        [
            ["Full model best", "Train", "0.2396", "0.8585", "0.7359", "0.7079"],
            ["Full model best", "Val", "0.3151", "0.6824", "0.4592", "0.5186"],
            ["Single-station ref", "Train", "0.3268", "0.7599", "0.5767", "0.5979"],
            ["Single-station ref", "Val", "0.3520", "0.5896", "0.2936", "0.4447"],
        ],
        columns=["Model", "Split", "MAE", "Corr", "R²", "Slope"],
    )
    add_table(s, metrics, 0.65, 1.22, 6.35, 1.68, 8.2, header_color=GREEN, widths=[2.2, 0.9, 0.95, 0.95, 0.95, 0.95])
    add_bullets(
        s,
        [
            "multi-station 结果相对 single-station 参考模型有明显提升。",
            "说明多台站空间信息对目标位置 PGA 估计有实际贡献。",
            "当前结果可作为摄像头位置参考 PGA 生成流程的模型基础。",
        ],
        7.35,
        1.35,
        5.10,
        1.22,
        size=10.8,
        bullet_color=GREEN,
    )
    add_image_fit(s, pga_asset("scatter_japan_overfit_pga15_cross_attention_overfitall_fixed_inputs_random_targets_best_train.png"), 0.70, 3.30, 5.85, 2.95)
    add_image_fit(s, pga_asset("scatter_japan_overfit_pga15_cross_attention_overfitall_fixed_inputs_random_targets_best_val.png"), 6.85, 3.30, 5.85, 2.95)
    add_pill(s, "Train", 2.85, 6.38, 1.40, 0.28, PANEL_GREEN, GREEN)
    add_pill(s, "Validation", 8.85, 6.38, 1.55, 0.28, PANEL_ORANGE, ORANGE)


def station_count_slide(d: Deck):
    s = d.slide("多台站数量与空间信息的作用", "输入台站数增加时，validation 误差大体下降", "Evidence", CYAN)
    add_image_fit(s, pga_asset("val_mae_by_station_count.png"), 0.60, 1.20, 5.75, 3.55)
    add_image_fit(s, pga_asset("attention_target_4_best_val_n12_case15.png"), 6.65, 1.20, 5.85, 3.55)
    buckets = pd.DataFrame(
        [
            ["2-3", 271, "0.3382", "0.6012"],
            ["4-5", 158, "0.3130", "0.7166"],
            ["6-10", 340, "0.3001", "0.7494"],
            ["11-15", 188, "0.2779", "0.7959"],
            ["16+", 435, "0.2698", "0.7607"],
        ],
        columns=["Input stations", "Targets", "MAE", "Corr"],
    )
    add_table(s, buckets, 0.82, 5.12, 5.25, 1.15, 8.2, header_color=CYAN, widths=[1.5, 1.0, 1.0, 1.0])
    add_card(s, 6.85, 5.10, 5.35, 1.10, "解释", "目标位置 query 会从多个输入台站读取信息；这支持把摄像头位置作为目标点查询。", CYAN, PANEL)


def elapsed_time_slide(d: Deck):
    s = d.slide("输入波形时间长度对 PGA 估计的影响", "rt44-rt51 按当前输入时长统计；Test* 对应本地 val split", "Realtime", CYAN)
    detail = rt48_elapsed_detail_table()
    summary = rt_elapsed_test_mae_table()
    add_table(
        s,
        detail,
        0.55,
        1.25,
        5.55,
        2.80,
        7.3,
        header_color=CYAN,
        widths=[0.75, 1.05, 1.05, 1.05, 1.05],
    )
    add_table(
        s,
        summary,
        6.45,
        1.25,
        6.45,
        2.80,
        6.3,
        header_color=BLUE,
        widths=[0.70, 0.76, 0.76, 0.76, 0.76, 0.76, 0.76, 0.76],
    )
    add_card(
        s,
        0.75,
        4.65,
        3.65,
        1.00,
        "主要趋势",
        "从 1s 到 5-40s，测试 MAE 明显下降，Corr 上升。",
        CYAN,
        PANEL,
        title_size=11.5,
        body_size=10.0,
    )
    add_card(
        s,
        4.85,
        4.65,
        3.65,
        1.00,
        "rt48 表现",
        "Test MAE 从 1s 的 0.3915 降到 40s 的 0.2589。",
        GREEN,
        PANEL,
        title_size=11.5,
        body_size=10.0,
    )
    add_card(
        s,
        8.95,
        4.65,
        3.65,
        1.00,
        "90s 现象",
        "90s 未继续稳定提升，说明更长窗口不必然更好。",
        ORANGE,
        PANEL,
        title_size=11.5,
        body_size=10.0,
    )
    add_pill(s, "左表：rt48 train/test 详细结果；右表：rt44-rt51 Test MAE* 汇总", 1.65, 6.35, 10.05, 0.32, PANEL_BLUE, BLUE)


def training_status_slide(d: Deck):
    s = d.slide("当前模型训练阶段性状态", "rt44-rt51 最新一组 PGA 结果；完整本地结果不含 rt47", "Status", PURPLE)
    rt_table = rt44_rt51_table()
    add_table(
        s,
        rt_table,
        0.35,
        1.18,
        12.65,
        4.05,
        6.8,
        header_color=PURPLE,
        widths=[0.70, 2.20, 0.88, 0.88, 0.88, 0.88, 0.88, 0.88],
    )
    add_card(
        s,
        0.85,
        5.70,
        11.65,
        0.68,
        "说明",
        "Test* 对应本地 eval 文件中的 val split；rt47 只有残缺 txt，没有完整 npz，因此未纳入表格。",
        ORANGE,
        PANEL_ORANGE,
        title_size=10.8,
        body_size=9.5,
    )


def calibration_workflow_slide(d: Deck):
    s = d.slide("如何获得摄像头标定系数", "摄像头位置没有台站记录，因此需要模型先估计该位置的参考 PGA", "Calibration", ORANGE)
    add_shape(s, 0.78, 1.38, 5.25, 3.25, PANEL_BLUE, LINE, radius=True)
    add_textbox(s, "历史地震事件", 1.05, 1.65, 4.70, 0.28, 14, BLUE, True, PP_ALIGN.CENTER, margin=0.0)
    station_points = [(1.55, 2.35), (2.40, 3.25), (3.38, 2.62), (4.85, 3.12)]
    for j, (x, y) in enumerate(station_points):
        add_shape(s, x, y, 0.22, 0.22, BLUE, BLUE, radius=True)
        add_textbox(s, f"S{j+1}", x - 0.08, y + 0.26, 0.38, 0.14, 7.5, BLUE, True, PP_ALIGN.CENTER, margin=0.0)
    add_shape(s, 3.95, 2.10, 0.42, 0.30, ORANGE, ORANGE, radius=True)
    add_textbox(s, "Camera", 3.53, 2.48, 1.20, 0.18, 8.5, ORANGE, True, PP_ALIGN.CENTER, margin=0.0)
    add_textbox(s, "没有同址强震台站", 3.18, 2.85, 1.85, 0.18, 9.5, RED, True, PP_ALIGN.CENTER, margin=0.0)
    add_line(s, 4.15, 2.40, 4.15, 2.82, RED, 1.2)

    add_arrow(s, 6.20, 2.98, 7.05, 2.98, CYAN, 1.8)
    add_card(s, 7.20, 2.20, 2.25, 1.25, "深度 PGA 模型", "输入台站波形与位置，查询摄像头位置", CYAN, PANEL, title_size=11.5, body_size=9.5)
    add_arrow(s, 9.60, 2.98, 10.35, 2.98, CYAN, 1.8)
    add_card(s, 10.55, 2.20, 2.05, 1.25, "参考 PGA", "PGA_camera,event", GREEN, PANEL_GREEN, title_size=11.5, body_size=10.0)

    add_shape(s, 0.90, 5.10, 11.60, 0.82, PANEL_ORANGE, LINE, radius=True)
    add_textbox(s, "有了多个历史事件的参考 PGA 和视频振动强度，就可以拟合：PGA_camera,event = k_camera × Motion_video,event", 1.15, 5.34, 11.10, 0.28, 13.6, NAVY, True, PP_ALIGN.CENTER, margin=0.0)


def traditional_comparison_slide(d: Deck):
    s = d.slide("与传统 PGA 方法的同口径对比", "比较对象要统一事件、目标位置、单位和评价指标", "Comparison", BLUE)
    add_card(s, 0.90, 1.30, 5.45, 1.05, "同一批事件", "使用相同地震事件、相同台站输入和相同目标摄像头位置。", BLUE, PANEL, number="1")
    add_card(s, 6.95, 1.30, 5.45, 1.05, "同一套标签/单位", "统一 PGA 单位、log 变换、水平/三分量定义和时间窗口。", CYAN, PANEL, number="2")
    add_card(s, 0.90, 2.80, 5.45, 1.05, "点位误差", "MAE、RMSE、Bias、Corr、Slope；按 PGA 强弱分桶。", GREEN, PANEL, number="3")
    add_card(s, 6.95, 2.80, 5.45, 1.05, "空间误差", "摄像头/目标点 residual map、距离分桶、近场/远场表现。", ORANGE, PANEL, number="4")
    add_shape(s, 1.10, 4.70, 11.05, 0.95, PANEL_BLUE, LINE, radius=True)
    add_textbox(s, "传统方法可以作为 baseline、物理先验或后处理校准参考；深度模型则提供从多台站波形到任意目标位置 PGA 的数据驱动估计。", 1.40, 4.95, 10.45, 0.30, 13.5, NAVY, True, PP_ALIGN.CENTER, margin=0.0)


def next_steps_slide(d: Deck):
    s = d.slide("后续工作安排", "围绕摄像头标定闭环推进，而不是只继续做模型分数", "Next", GREEN)
    tracks = [
        ("1. 生成参考 PGA", "用多台站模型批量查询历史事件的摄像头位置 PGA。", BLUE),
        ("2. 汇总视频指标", "把每个历史事件的视频振动曲线转成可拟合强度指标。", CYAN),
        ("3. 拟合标定系数", "按摄像头位置拟合 k_camera，并做事件级留出验证。", ORANGE),
        ("4. 同口径对比", "与传统 PGA 方法比较点位误差和空间残差。", GREEN),
    ]
    for i, (title, body, color) in enumerate(tracks):
        add_card(s, 0.95, 1.25 + i * 1.15, 11.45, 0.80, title, body, color, PANEL)
    add_pill(s, "最终交付：摄像头位置 PGA 估计流程 + 标定系数 + 与传统方法的同口径对比结果", 1.15, 6.15, 11.05, 0.36, PANEL_GREEN, GREEN)


def main():
    d = Deck()
    title_slide(d)
    workflow_slide(d)
    task_split_slide(d)
    calibration_factor_slide(d)
    calibration_workflow_slide(d)
    model_io_slide(d)
    architecture_slide(d)
    data_training_slide(d)
    data_spatial_distribution_slide(d)
    results_slide(d)
    elapsed_time_slide(d)
    training_status_slide(d)
    traditional_comparison_slide(d)
    next_steps_slide(d)
    d.save(PPTX_PATH)
    print(f"PPTX: {PPTX_PATH}")


if __name__ == "__main__":
    main()
