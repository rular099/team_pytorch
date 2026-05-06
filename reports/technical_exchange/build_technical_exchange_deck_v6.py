from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

from build_technical_exchange_deck_v5 import (
    ASSET_DIR,
    TABLE_DIR,
    BG,
    BLUE,
    CYAN,
    GREEN,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    PANEL,
    PANEL_BLUE,
    PANEL_GREEN,
    PANEL_ORANGE,
    PANEL_RED,
    PURPLE,
    RED,
    SLIDE_H,
    SLIDE_W,
    TEXT,
    TEAL,
    WHITE,
    Deck,
    add_bullets,
    add_card,
    add_image_fit,
    add_kpi,
    add_line,
    add_pill,
    add_shape,
    add_table,
    add_textbox,
)

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "diting_team_graph_pga_technical_exchange_v6.pptx"


def path_asset(name: str) -> Path:
    p = ASSET_DIR / name
    if not p.exists():
        raise FileNotFoundError(f"Missing asset: {p}")
    return p


def read_table(name: str) -> pd.DataFrame:
    p = TABLE_DIR / name
    if not p.exists():
        raise FileNotFoundError(f"Missing table: {p}")
    return pd.read_csv(p)


def title_slide(d: Deck):
    s = d.prs.slides.add_slide(d.blank)
    d.page_no += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    add_shape(s, 0, 0, SLIDE_W, 0.12, CYAN, CYAN, radius=False)
    add_shape(s, 8.75, 0.72, 3.95, 5.95, PANEL_BLUE, PANEL_BLUE, radius=True)
    add_shape(s, 9.20, 1.15, 3.05, 1.00, WHITE, LINE, radius=True)
    add_shape(s, 9.20, 2.45, 3.05, 1.00, WHITE, LINE, radius=True)
    add_shape(s, 9.20, 3.75, 3.05, 1.00, WHITE, LINE, radius=True)
    add_textbox(s, "组会讨论 · 技术进展同步", 0.72, 1.08, 5.0, 0.34, 15.5, CYAN, True, margin=0.0)
    add_textbox(s, "DiTing 表征接入\n多台站 PGA / Event 预测", 0.68, 1.62, 7.7, 1.35, 32, NAVY, True, margin=0.0)
    add_textbox(s, "Transformer readout 与 graph prior-residual 的实验诊断", 0.72, 3.18, 7.45, 0.35, 15, TEXT, False, margin=0.0)
    add_shape(s, 0.72, 5.78, 2.10, 0.36, CYAN, CYAN, radius=True)
    add_textbox(s, "v6 讨论版", 0.72, 5.81, 2.10, 0.30, 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0.0)
    add_textbox(s, "2026-05-02", 0.72, 6.28, 2.4, 0.24, 10, MUTED, False, margin=0.0)
    add_textbox(s, "Task setup", 9.52, 1.35, 2.4, 0.26, 11.5, BLUE, True, PP_ALIGN.CENTER, margin=0.0)
    add_textbox(s, "Cross-attention route", 9.52, 2.65, 2.4, 0.26, 11.5, CYAN, True, PP_ALIGN.CENTER, margin=0.0)
    add_textbox(s, "Graph exploration", 9.52, 3.95, 2.4, 0.26, 11.5, GREEN, True, PP_ALIGN.CENTER, margin=0.0)
    add_textbox(s, "目标：同步当前实验进展，讨论 readout、graph 与数据质量的下一步实验设计。", 0.72, 4.38, 7.2, 0.48, 15.0, NAVY, True, margin=0.0)
    return s


def rename_train_eval(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={
        "Val MAE": "Eval MAE",
        "Val Corr": "Eval Corr",
        "Val R2": "Eval R2",
        "Val slope": "Eval slope",
    })


def main():
    all_exp = read_table("all_experiments_summary.csv")
    transformer = read_table("transformer_ablation.csv")
    cross = read_table("cross_attention.csv")
    graph = read_table("graph_results.csv")
    variance = read_table("variance_summary.csv")
    single = read_table("single_station_metrics.csv")
    model_scale = read_table("model_scale_params.csv")
    loss_design = read_table("loss_design.csv")

    # Display tables
    single_val = single[single["Split"].eq("val")].copy()
    single_pga = single_val[single_val["Task"].eq("pga")][["Model", "Samples", "MAE", "Corr"]].copy()
    single_all = single_val[["Model", "Task", "Samples", "MAE", "Corr"]].copy()

    transformer_disp = rename_train_eval(transformer[[
        "Method", "Train MAE", "Train Corr", "Train R2", "Train slope",
        "Val MAE", "Val Corr", "Val R2", "Val slope"
    ]].copy())

    cross_disp = rename_train_eval(cross[[
        "Experiment", "Train MAE", "Train Corr", "Train R2",
        "Val MAE", "Val Corr", "Val R2"
    ]].copy())
    cross_disp["Experiment"] = cross_disp["Experiment"].replace({
        "pga15_cross_overfit32": "overfit32 random",
        "pga15_cross_overfit128": "overfit128 random",
        "fixed_inputs_targets": "fixed inputs/targets",
        "input_targets": "input stations as targets",
        "event_pga_cross_first_inputs": "event+pga first inputs",
    })

    graph_disp = rename_train_eval(graph[[
        "Experiment", "Train MAE", "Train Corr", "Train R2",
        "Val MAE", "Val Corr", "Val R2"
    ]].copy())
    graph_disp["Experiment"] = graph_disp["Experiment"].replace({
        "old graph first-inputs": "old graph first-inputs",
        "prior-residual first-inputs": "prior-residual first-inputs",
        "prior-residual random-inputs": "prior-residual random-inputs",
        "graph exp1 same-station": "exp1 same-station",
        "graph exp2 multi-target": "exp2 multi-target",
        "graph exp3 holdout": "exp3 holdout",
    })

    old_graph = graph_disp[graph_disp["Experiment"].isin([
        "old graph first-inputs", "exp1 same-station", "exp2 multi-target", "exp3 holdout"
    ])].copy()
    prior_graph = graph_disp[graph_disp["Experiment"].isin([
        "old graph first-inputs", "prior-residual first-inputs", "prior-residual random-inputs"
    ])].copy()

    d = Deck()
    title_slide(d)

    # 2 Agenda / current work map, no conclusion
    s = d.slide("今天汇报的内容", "按工作流展开：任务设置 → 表征与 readout → 两条路线 → 数据风险 → 下一步", "Agenda", CYAN)
    items = [
        ("1. 任务与数据边界", "从 DiTing 单台站表征到多台站 event / target PGA。", BLUE),
        ("2. Single-station 阶段", "说明 station encoder / adapter 的预训练和基本指标。", CYAN),
        ("3. Transformer / cross-attention 路线", "从原始 query_transformer 到显式 target-to-station readout。", ORANGE),
        ("4. Graph 路线", "从 old graph message passing 到 prior-residual 探索。", GREEN),
        ("5. P-pick 与下一步", "数据质量风险、后续主要实验计划。", RED),
    ]
    for i, (t, b, c) in enumerate(items):
        add_card(s, 0.95, 1.25 + i * 0.98, 11.45, 0.68, t, b, c, PANEL, title_size=11.8, body_size=10.2)

    # 3 Task setup
    s = d.slide("任务转换：从单台站 DiTing 到多台站 event/PGA", "当前任务从局部 station-level 预测变为空间场 readout", "Task", BLUE)
    add_image_fit(s, path_asset("task_setup_multistation.png"), 0.65, 1.28, 8.25, 4.65)
    add_card(s, 9.25, 1.30, 3.35, 1.05, "输入", "N 个 station 波形 + 坐标 + station_valid mask", BLUE, PANEL)
    add_card(s, 9.25, 2.65, 3.35, 1.05, "查询", "event query；PGA target query", CYAN, PANEL)
    add_card(s, 9.25, 4.00, 3.35, 1.05, "输出", "event-level mag/loc；target-level PGA", GREEN, PANEL)
    add_textbox(s, "核心变化：target station 与输入 station 不一定重合，PGA 需要空间传播建模。", 9.35, 5.55, 3.1, 0.42, 11.2, TEXT, True, margin=0.0)

    # 4 Data and overfit boundary
    s = d.slide("数据覆盖与 overfit 诊断边界", "当前是结构诊断阶段，不把小验证集结果解读为最终泛化结论", "Data", BLUE)
    add_image_fit(s, path_asset("data_overview.png"), 0.55, 1.25, 6.25, 4.75)
    add_image_fit(s, path_asset("sample_event_geometry.png"), 6.95, 1.25, 5.85, 3.15)
    add_card(s, 7.05, 4.70, 2.75, 1.02, "关注", "train 是否可记忆；eval 是否辅助监控；输出是否非平凡。", CYAN, PANEL)
    add_card(s, 10.05, 4.70, 2.75, 1.02, "避免", "只看 loss；把小样本 eval 过度解释为泛化。", ORANGE, PANEL)
    add_pill(s, "指标组合：MAE · Corr · R² · slope · pred std / target std", 7.18, 5.98, 5.35, 0.32, PANEL_BLUE, BLUE)

    # 5 Implementation overview
    s = d.slide("实现级模型总览：共用 encoder + 不同 readout", "后续结果按 readout route 分组解释", "Model", CYAN)
    add_image_fit(s, path_asset("architecture_implementation_detail.png"), 0.45, 1.18, 12.45, 5.52)

    # 6 Model scale and loss
    s = d.slide("模型规模与训练目标", "当前主要训练 adapter / readout / task heads；DiTing encoder 冻结", "Setup", CYAN)
    scale = model_scale[["Component", "Total params", "Trainable params", "Status"]].copy()
    add_table(s, scale, 0.55, 1.28, 6.1, 2.25, 8.2, widths=[2.4, 1.2, 1.4, 1.2])
    loss = loss_design[["Stage", "Tasks", "Weights", "Implementation note"]].copy()
    add_table(s, loss, 6.95, 1.28, 5.85, 2.25, 7.2, header_color=TEAL, widths=[1.5, 1.3, 1.1, 2.5])
    add_card(s, 0.75, 4.15, 3.75, 1.05, "参数", "TEAM cross-attn 约 28M 可训练；graph prior-residual 约 38M。", BLUE, PANEL)
    add_card(s, 4.82, 4.15, 3.75, 1.05, "Loss", "Huber point prediction；均值解也可能降低 loss。", ORANGE, PANEL)
    add_card(s, 8.88, 4.15, 3.75, 1.05, "评估", "loss 需要和 corr、slope、pred std 一起看。", RED, PANEL)

    # 7 Route map: two independent routes
    s = d.slide("两条 readout 路线的实验组织", "Transformer→cross-attention 与 old graph→prior-residual 是两条并行路线", "Routes", CYAN)
    # Left route
    add_shape(s, 0.85, 1.45, 5.65, 4.50, PANEL, LINE, radius=True)
    add_textbox(s, "路线 A：Transformer / Cross-attention", 1.05, 1.68, 5.20, 0.28, 14, BLUE, True, PP_ALIGN.CENTER, margin=0.0)
    add_card(s, 1.25, 2.30, 4.75, 0.85, "Original query_transformer", "event / station / PGA tokens 共同进入 full Transformer。", ORANGE, PANEL_ORANGE)
    add_line(s, 3.62, 3.15, 3.62, 3.70, ORANGE, 1.4)
    add_card(s, 1.25, 3.72, 4.75, 0.85, "PGA / event cross-attention", "target/event query 显式读取 station tokens。", BLUE, PANEL_BLUE)
    add_textbox(s, "后续主要沿这条路线推进。", 1.55, 5.10, 4.1, 0.25, 11.5, BLUE, True, PP_ALIGN.CENTER, margin=0.0)
    # Right route
    add_shape(s, 6.85, 1.45, 5.65, 4.50, PANEL, LINE, radius=True)
    add_textbox(s, "路线 B：Graph 探索", 7.05, 1.68, 5.20, 0.28, 14, GREEN, True, PP_ALIGN.CENTER, margin=0.0)
    add_card(s, 7.25, 2.30, 4.75, 0.85, "Old graph message passing", "station-to-target graph readout baseline。", RED, PANEL_RED)
    add_line(s, 9.62, 3.15, 9.62, 3.70, GREEN, 1.4)
    add_card(s, 7.25, 3.72, 4.75, 0.85, "Graph prior-residual", "single-station prior + distance baseline + residual。", GREEN, PANEL_GREEN)
    add_textbox(s, "作为探索路线继续保留。", 7.55, 5.10, 4.1, 0.25, 11.5, GREEN, True, PP_ALIGN.CENTER, margin=0.0)

    # 8 Single station setup
    s = d.slide("Single-station 阶段：设置与评估口径", "这些指标来自各实验目录中的 single-station pretrain/readout 评估", "Single-station", CYAN)
    add_card(s, 0.85, 1.40, 3.65, 1.20, "输入", "单台站三分量波形 + 台站信息；共享 frozen DiTing encoder。", BLUE, PANEL)
    add_card(s, 4.85, 1.40, 3.65, 1.20, "训练目标", "mag / epidist / PGA；Huber loss 加权 0.3 / 0.3 / 0.4。", CYAN, PANEL)
    add_card(s, 8.85, 1.40, 3.65, 1.20, "用途", "检查 waveform station representation 是否含有任务相关信号。", GREEN, PANEL)
    add_table(s, loss_design[loss_design["Stage"].eq("single-station pretrain")][["Stage", "Tasks", "Loss", "Weights", "Implementation note"]], 0.95, 3.18, 11.45, 0.75, 8.8, header_color=CYAN, widths=[1.8, 1.8, 1.5, 1.5, 3.0])
    add_textbox(s, "说明：表中的 Model 名称如 “Cross overfit128 / Graph prior first-inputs” 是对应 checkpoint/run 的标签；本页指标本身是 single-station 分支的 train/eval 结果，不是 full multi-station PGA 结果。", 1.05, 4.65, 11.2, 0.55, 12.0, TEXT, True, PP_ALIGN.CENTER, margin=0.0)

    # 9 Single station metrics
    s = d.slide("Single-station model 指标", "用于观察 station waveform 表征的基本可用性", "Single-station", CYAN)
    add_image_fit(s, path_asset("single_station_val_metrics.png"), 0.55, 1.18, 6.45, 4.90)
    add_table(s, single_all, 7.25, 1.25, 5.35, 3.65, 7.0, header_color=CYAN, widths=[2.2, 0.9, 0.9, 0.9, 0.9])
    add_textbox(s, "右表为 single-station val 指标；左图按任务展示 Corr 和 MAE。后续 full model 的问题需要和这些单台站指标分开讨论。", 7.35, 5.35, 5.15, 0.48, 11.5, MUTED, False, margin=0.0)

    # 10 Original TEAM readout design
    s = d.slide("路线 A-1：原始 TEAM / query_transformer readout", "统一 token 模型的实现方式", "Transformer", ORANGE)
    add_image_fit(s, path_asset("architecture_team.png"), 0.65, 1.25, 6.15, 4.45)
    add_bullets(s, [
        "输入 token：event token、station tokens、PGA target tokens。",
        "PGA query = target coordinate embedding + learned PGA query token。",
        "PGA key 被 mask，PGA token 主要作为 query-only token。",
        "这个设计让模型自己学习 event/station/target 的 readout 路径。",
    ], 7.25, 1.55, 5.10, 2.20, size=12.5, bullet_color=ORANGE)
    add_card(s, 7.20, 4.55, 5.20, 0.92, "讨论点", "如果 readout 路由学习不稳定，loss 下降也可能对应均值/常数解。", ORANGE, PANEL_ORANGE)

    # 11 Ablation settings larger
    s = d.slide("原始 readout ablation：实验设置", "先说明每个 ablation 在检查什么", "Transformer", ORANGE)
    settings = pd.DataFrame([
        ["query_transformer", "PGA query in full self-attention", "原始 readout；检查 full Transformer 是否能自动学 route"],
        ["mask_batch1", "batch=1 mask sanity", "排查跨 batch 泄露或 mask 维度问题"],
        ["query_no_transformer", "target coord + learned query only", "负对照：不读取 station 信息，只看 target/query shortcut"],
        ["direct_station", "pooled station embedding readout", "检查 station feature 和 PGA head 是否可读出信号"],
    ], columns=["Method", "Setting", "Purpose"])
    add_table(s, settings, 0.75, 1.45, 11.85, 3.15, 11.0, header_color=ORANGE, widths=[1.6, 2.4, 4.0])
    add_textbox(s, "这页只讲设置；下一页单独看指标，避免图片和表格过小。", 1.05, 5.35, 11.1, 0.35, 12.0, MUTED, False, PP_ALIGN.CENTER, margin=0.0)

    # 12 Ablation metrics large table
    s = d.slide("原始 readout ablation：指标结果", "train 与 eval、corr 与 slope 一起看", "Transformer", ORANGE)
    add_table(s, transformer_disp, 0.35, 1.25, 12.65, 3.15, 7.4, header_color=ORANGE, widths=[1.8, 1, 1, 1, 1, 1, 1, 1, 1])
    add_bullets(s, [
        "query_transformer 和 mask_batch1 的 slope 接近 0，说明输出变化被压扁。",
        "direct_station 可作为 station feature/head 是否可用的对照，不等价于任意 target 空间传播。",
        "query_no_transformer 是 target/query-only 对照，需要和读取 station 的模型分开解释。",
    ], 0.85, 4.85, 11.55, 1.25, size=12.5, bullet_color=ORANGE)

    # 13 Variance/diagnostics
    s = d.slide("原始 readout 诊断：输出方差与均值解风险", "loss 下降不能单独排除常数化", "Transformer", ORANGE)
    add_image_fit(s, path_asset("variance_compression.png"), 0.75, 1.25, 7.25, 4.20)
    add_table(s, variance[variance["Experiment"].isin(["query_transformer", "direct_station", "cross fixed targets"])], 8.35, 1.45, 4.05, 1.55, 8.0, header_color=ORANGE, widths=[1.8, 1.0, 1.0])
    add_card(s, 8.35, 3.50, 4.05, 1.15, "观察口径", "均值解也能降低 Huber loss；因此后续统一同时记录 corr、slope、pred std。", RED, PANEL_RED)
    add_textbox(s, "注：0.11111 attention 诊断 bug 已修，已区分 key-type mask 与 padding 后 effective key mask。", 8.45, 5.20, 3.85, 0.38, 10.8, MUTED, False, margin=0.0)

    # 14 Cross design
    s = d.slide("路线 A-2：Cross-attention readout", "把 target/event query 读取 station tokens 的路径显式化", "Cross", BLUE)
    add_image_fit(s, path_asset("architecture_cross_attention.png"), 0.55, 1.28, 6.55, 4.65)
    add_card(s, 7.40, 1.45, 4.95, 0.95, "Q", "target coordinate embedding + learned PGA token", BLUE, PANEL)
    add_card(s, 7.40, 2.65, 4.95, 0.95, "K / V", "station_feature_emb；mask = ~station_valid", CYAN, PANEL)
    add_card(s, 7.40, 3.85, 4.95, 0.95, "Readout", "输出不加 query residual，更直接依赖 station tokens", GREEN, PANEL)
    add_textbox(s, "动机：减少 full self-attention 中 event/PGA token 路由自发现的难度。", 7.45, 5.35, 4.75, 0.35, 11.5, TEXT, True, margin=0.0)

    # 15 Cross settings and results balanced
    s = d.slide("Cross-attention 实验结果", "train 和 eval 同样展示；eval 仅作为小样本监控，不作最终泛化判断", "Cross", BLUE)
    add_table(s, cross_disp, 0.35, 1.20, 12.65, 2.65, 7.8, header_color=BLUE, widths=[2.3, 1, 1, 1, 1, 1, 1])
    add_kpi(s, "overfit128", "0.843 / 0.618", "Train/Eval Corr", 0.80, 4.35, 2.0, 0.95, BLUE)
    add_kpi(s, "fixed targets", "0.968 / 0.621", "Train/Eval Corr", 3.15, 4.35, 2.0, 0.95, CYAN)
    add_kpi(s, "input targets", "0.991 / 0.641", "Train/Eval Corr", 5.50, 4.35, 2.0, 0.95, GREEN)
    add_bullets(s, [
        "当前更关注小样本 overfit / sanity 是否能学到非平凡映射。",
        "fixed/input target 设置用于检查 route 是否通畅；random target 设置用于辅助观察空间读出。",
        "event+pga 联合训练仍需单独拆 event readout 和 PGA context 的设计。",
    ], 7.95, 4.25, 4.75, 1.42, size=11.2, bullet_color=BLUE)

    # 16 Cross next route, not problem wording
    s = d.slide("Cross-attention 路线：后续主要推进方向", "目前这条路线效果更好，下一步围绕 readout 与 event context 做消融", "Cross", BLUE)
    add_card(s, 0.85, 1.35, 3.65, 1.15, "Readout 消融", "query residual、target position encoding、station mask、target 数量。", BLUE, PANEL)
    add_card(s, 4.85, 1.35, 3.65, 1.15, "Distance / relative feature", "relative coordinate MLP、distance bias、距离分桶或连续 bias。", CYAN, PANEL)
    add_card(s, 8.85, 1.35, 3.65, 1.15, "Event context", "event-only cross-attn；mag/loc context gate；分阶段训练。", GREEN, PANEL)
    add_image_fit(s, path_asset("cross_train_val_detail.png"), 1.10, 3.05, 5.55, 2.75)
    add_textbox(s, "说明：当前 eval 弱不单独作为问题，因为这些实验仍是小样本/overfit 诊断；重点是 train/sanity 是否说明 readout route 能工作。", 7.15, 3.35, 5.0, 1.05, 13, TEXT, True, margin=0.0)
    add_pill(s, "主线：继续沿 cross-attention 做结构消融与更大 split 验证", 7.20, 4.90, 4.85, 0.34, PANEL_BLUE, BLUE)

    # 17 Old graph no image
    s = d.slide("路线 B-1：Old graph message passing baseline", "这里不放图，直接看 4 个 baseline 设置和指标", "Graph", RED)
    add_table(s, old_graph, 0.45, 1.35, 12.45, 2.65, 8.4, header_color=RED, widths=[2.2, 1, 1, 1, 1, 1, 1])
    add_bullets(s, [
        "old graph first-inputs、single-input、multi-target、holdout 等设置均没有稳定表现出非平凡空间 readout。",
        "这说明仅加入 graph message passing 还不够，需要进一步加入 station prior 或距离传播约束。",
        "Graph 路线目前定位为探索路线，不作为主线替代 cross-attention。",
    ], 0.85, 4.55, 11.6, 1.25, size=12.5, bullet_color=RED)

    # 18 Prior residual design
    s = d.slide("路线 B-2：Graph prior-residual 设计", "把 single-station PGA prior 和距离 baseline 显式接入 graph readout", "Graph", GREEN)
    add_image_fit(s, path_asset("architecture_graph_prior_residual.png") if (ASSET_DIR / "architecture_graph_prior_residual.png").exists() else path_asset("architecture_graph_detail.png"), 0.55, 1.18, 7.05, 4.88)
    add_card(s, 7.90, 1.32, 4.55, 0.90, "1. Station prior", "station_pga_prior_head 从 single-station pretrain 加载。", CYAN, PANEL)
    add_card(s, 7.90, 2.45, 4.55, 0.90, "2. Distance baseline", "输入台站 prior 的距离加权平均，形成传播初值。", BLUE, PANEL)
    add_card(s, 7.90, 3.58, 4.55, 0.90, "3. Learned residual", "GraphPGAReadout 用 edge features 学修正项。", GREEN, PANEL)
    add_shape(s, 7.90, 5.05, 4.55, 0.55, PANEL_GREEN, LINE, radius=True)
    add_textbox(s, "final PGA = distance baseline + learned residual", 8.08, 5.20, 4.20, 0.22, 13, GREEN, True, PP_ALIGN.CENTER, margin=0.0)

    # 19 Prior residual results balanced no image
    s = d.slide("Graph prior-residual 结果", "train 和 eval 同样展示；这里主要比较 old graph 与 prior-residual 的差异", "Graph", GREEN)
    add_table(s, prior_graph, 0.45, 1.35, 12.45, 2.00, 8.4, header_color=GREEN, widths=[2.5, 1, 1, 1, 1, 1, 1])
    vgraph = variance[variance["Experiment"].isin(["graph old", "graph prior", "graph random prior"])].copy()
    add_table(s, vgraph, 1.00, 3.95, 5.80, 1.25, 8.3, header_color=GREEN, widths=[2.0, 1.1, 1.1])
    add_bullets(s, [
        "prior-residual 相比 old graph 有明显改善，但整体仍作为探索方向。",
        "后续如果继续做 graph，重点应是 baseline-only / prior-only / residual-only 的消融。",
        "当前主线仍建议放在 cross-attention。",
    ], 7.25, 3.85, 5.15, 1.45, size=11.8, bullet_color=GREEN)

    # 20 Route comparison revised
    s = d.slide("两条路线的阶段性定位", "Cross-attention 作为后续主线；graph prior-residual 继续探索", "Synthesis", CYAN)
    comp = pd.DataFrame([
        ["Cross-attention", "目前效果更好；route sanity 清楚；适合继续做消融", "主线推进：readout、distance bias、event context、更大 split"],
        ["Graph prior-residual", "引入 prior/baseline 后有改善，具有可解释性", "探索保留：做 prior/baseline/residual/kNN 消融"],
    ], columns=["路线", "当前定位", "后续安排"])
    add_table(s, comp, 0.65, 1.45, 12.05, 1.75, 10.0, header_color=CYAN, widths=[1.6, 3.6, 3.2])
    add_card(s, 1.00, 4.05, 5.35, 1.25, "近期主线", "沿 cross-attention 路线扩大实验，优先把 readout、target/query、event context 的消融做清楚。", BLUE, PANEL)
    add_card(s, 6.95, 4.05, 5.35, 1.25, "Graph 位置", "保持探索，不做融合方案；先确认 prior-residual 的每个组成是否真的贡献有效。", GREEN, PANEL)

    # 21 P-pick metadata
    s = d.slide("P-pick 和元数据质量风险", "P-pick 会影响波形窗口、station SNR 和 feature/label 对齐", "Data Risk", RED)
    add_image_fit(s, path_asset("ppick_metadata_stats.png"), 0.65, 1.20, 7.05, 4.85)
    add_bullets(s, [
        "当前流程：走时曲线粗定位 + STA/LTA 搜索窗 refine。",
        "需要检查 refined pick 与 travel-time pick 的系统偏差。",
        "低 SNR 或错误 pick 会影响 single-station prior 和 full-model readout。",
        "后续需要建立人工抽查样本集，而不是只看统计图。",
    ], 8.15, 1.55, 4.35, 2.15, size=12.3, bullet_color=RED)
    add_card(s, 8.10, 4.55, 4.35, 0.95, "注意", "部分 p_picks 是 aligned/global sample，超出当前 100s 训练窗口，需要回到原始 aligned 记录抽查。", RED, PANEL_RED)

    # 22 waveform full slide
    s = d.slide("P-pick 波形抽样检查", "单独放大展示，便于现场讨论起跳点是否合理", "Data Risk", RED)
    add_image_fit(s, path_asset("ppick_audit_waveforms.png"), 0.30, 1.05, 12.75, 5.75)
    add_textbox(s, "红虚线为当前存储 P-pick；这里只展示 pick 落在当前 100s HDF5 窗口内的样本。", 0.60, 6.86, 12.10, 0.22, 9.8, MUTED, False, PP_ALIGN.CENTER, margin=0.0)

    # 23 Next experiments
    s = d.slide("下一步实验计划", "组会后优先把主线实验和数据质检排清楚", "Plan", GREEN)
    tracks = [
        ("P0 数据质检", "P-pick 人工抽查；低 SNR / STA-LTA 失败类型统计；必要时剔除或降权。", RED),
        ("P1 Cross-attention 主线", "readout route、target query、relative coordinate / distance bias、event context gate 消融。", BLUE),
        ("P2 扩大验证", "多 seed / 多 split；固定记录 train/eval MAE、Corr、R²、slope、pred std。", CYAN),
        ("P3 Graph 探索", "baseline-only、prior-only、residual-only、prior+baseline+residual、kNN / distance power p。", GREEN),
    ]
    for i, (t, b, c) in enumerate(tracks):
        add_card(s, 0.90, 1.35 + i * 1.23, 11.55, 0.88, t, b, c, PANEL, title_size=12.5, body_size=10.5)
    add_pill(s, "优先级：Cross-attention 主线 + P-pick 质检；Graph prior-residual 保持探索", 1.20, 6.25, 10.90, 0.36, PANEL_GREEN, GREEN)

    # 24 Final conclusion moved to end
    s = d.slide("阶段性总结", "把本次讨论收束到目前工作判断和后续路线", "Summary", CYAN)
    add_card(s, 0.85, 1.30, 3.65, 1.20, "1. 已完成", "搭建了 DiTing station 表征接入多台站 event/PGA 的实验框架，并完成多组 overfit/sanity 诊断。", BLUE, PANEL)
    add_card(s, 4.85, 1.30, 3.65, 1.20, "2. Single-station", "单台站分支在 mag / epidist / PGA 上有可观相关性，说明 station 表征值得继续使用。", CYAN, PANEL)
    add_card(s, 8.85, 1.30, 3.65, 1.20, "3. Readout", "原始 query_transformer 容易输出压扁；cross-attention 显式 readout 目前效果更好。", ORANGE, PANEL)
    add_card(s, 0.85, 3.10, 3.65, 1.20, "4. Graph", "old graph 不够；prior-residual 有改善，但当前定位为探索路线。", GREEN, PANEL)
    add_card(s, 4.85, 3.10, 3.65, 1.20, "5. 数据", "P-pick 需要人工抽查闭环，避免数据问题与模型问题混在一起。", RED, PANEL)
    add_card(s, 8.85, 3.10, 3.65, 1.20, "6. 后续", "主线沿 cross-attention 做消融和更大 split；graph prior-residual 继续补消融。", PURPLE, PANEL)
    add_shape(s, 0.95, 5.45, 11.35, 0.62, PANEL_BLUE, LINE, radius=True)
    add_textbox(s, "当前工作重点：把 cross-attention 主线做扎实，同时完成 P-pick 质检和 graph prior-residual 消融。", 1.18, 5.62, 10.90, 0.24, 13.5, NAVY, True, PP_ALIGN.CENTER, margin=0.0)

    d.save(PPTX_PATH)
    print(f"PPTX: {PPTX_PATH}")


if __name__ == "__main__":
    main()
